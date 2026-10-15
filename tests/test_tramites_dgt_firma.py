from docx import Document

from services.tramites_dgt_service import TramitesDgtService


def test_normaliza_indices_de_firma_del_contrato_en_plantilla_antigua(tmp_path):
    ruta = tmp_path / "contrato.docx"
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).paragraphs[0].add_run("Vendedor")
    table.cell(0, 0).add_paragraph("[[s|1                          ]]")
    table.cell(0, 1).paragraphs[0].add_run("Comprador")
    table.cell(0, 1).add_paragraph("[[s|2                          ]]")
    document.save(ruta)

    TramitesDgtService._asegurar_etiquetas_firma(ruta, "contrato_compraventa")

    corregido = Document(ruta)
    vendedor, comprador = corregido.tables[-1].rows[-1].cells[:2]
    texto_vendedor = "\n".join(p.text for p in vendedor.paragraphs)
    texto_comprador = "\n".join(p.text for p in comprador.paragraphs)
    assert "[[s|0" in texto_vendedor
    assert "[[s|1" in texto_comprador
    assert "[[s|2" not in texto_vendedor + texto_comprador


class _RepositorioFirma:
    def __init__(self):
        self.expediente = {
            "id": "exp-1",
            "firma_evidencia": {
                "solicitudes": [
                    {
                        "request_id": "firma-1",
                        "tipo_documento": "contrato_compraventa",
                        "firmantes_esperados": [
                            "vendedor@example.com",
                            "comprador@example.com",
                        ],
                    },
                ]
            },
        }

    def get_expediente(self, expediente_id):
        return self.expediente if expediente_id == "exp-1" else None

    def upsert_expediente(self, expediente):
        self.expediente = expediente
        return expediente["id"]


class _FirmaParcial:
    def consultar(self, _request_id):
        return {
            # Simula una respuesta incoherente del proveedor: documento
            # finalizado, pero falta uno de los firmantes que envio DGT.
            "status": "signed",
            "firmantes": [
                {"email": "vendedor@example.com", "needs_to_sign": True, "signed": True},
            ],
        }

    def descargar_evidencias(self, *_args, **_kwargs):
        raise AssertionError("No debe descargar evidencias con una firma pendiente")


def test_expediente_permanece_parcial_mientras_falte_un_firmante():
    repository = _RepositorioFirma()
    service = TramitesDgtService(repository=repository, firma_client=_FirmaParcial())

    resultado = service.actualizar_estado_firma("exp-1")

    assert resultado["estado"] == "parcialmente_firmado"
    assert repository.expediente["firma_estado"] == "parcialmente_firmado"
    solicitud = repository.expediente["firma_evidencia"]["solicitudes"][0]
    assert solicitud["firmantes_pendientes"] == ["comprador@example.com"]


class _RepositorioEnvio:
    def __init__(self):
        self.expediente = {
            "id": "exp-2",
            "referencia": "DGT-2026-0002",
            "firma_estado": "pendiente",
            "vendedor_payload": {"email": "vendedor@example.com"},
            "comprador_payload": {"email": "comprador@example.com"},
        }

    def get_expediente(self, expediente_id):
        return self.expediente if expediente_id == "exp-2" else None

    def listar_documentos_generados(self, _expediente_id):
        return [
            {
                "id": "doc-1",
                "tipo_documento": "contrato_compraventa",
                "titulo": "Contrato",
                "ruta_pdf": "contrato.pdf",
            }
        ]

    def upsert_expediente(self, expediente):
        self.expediente = expediente
        return expediente["id"]


class _FirmaEnvio:
    def enviar_documento(self, **_kwargs):
        return {"uuid": "firma-2", "status": "sent"}


def test_envio_guarda_los_firmantes_esperados_del_contrato():
    repository = _RepositorioEnvio()
    service = TramitesDgtService(repository=repository, firma_client=_FirmaEnvio())

    service.enviar_a_firma("exp-2")

    solicitud = repository.expediente["firma_evidencia"]["solicitudes"][0]
    assert solicitud["firmantes_esperados"] == [
        "vendedor@example.com",
        "comprador@example.com",
    ]
