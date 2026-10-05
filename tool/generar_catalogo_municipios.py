"""Genera el catalogo municipal incluido en el backend desde el Excel del INE."""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import load_workbook


CCAA_POR_PROVINCIA = {
    "01": "ES21", "02": "ES42", "03": "ES52", "04": "ES61",
    "05": "ES41", "06": "ES43", "07": "ES53", "08": "ES51",
    "09": "ES41", "10": "ES43", "11": "ES61", "12": "ES52",
    "13": "ES42", "14": "ES61", "15": "ES11", "16": "ES42",
    "17": "ES51", "18": "ES61", "19": "ES42", "20": "ES21",
    "21": "ES61", "22": "ES24", "23": "ES61", "24": "ES41",
    "25": "ES51", "26": "ES23", "27": "ES11", "28": "ES30",
    "29": "ES61", "30": "ES62", "31": "ES22", "32": "ES11",
    "33": "ES12", "34": "ES41", "35": "ES70", "36": "ES11",
    "37": "ES41", "38": "ES70", "39": "ES13", "40": "ES41",
    "41": "ES61", "42": "ES41", "43": "ES51", "44": "ES24",
    "45": "ES42", "46": "ES52", "47": "ES41", "48": "ES21",
    "49": "ES41", "50": "ES24", "51": "ES63", "52": "ES64",
}


def generar(origen: Path, destino: Path) -> int:
    workbook = load_workbook(origen, read_only=True, data_only=True)
    rows: list[tuple[str, str, str, str]] = []
    for sheet in workbook.worksheets:
        provincia = str(sheet["A2"].value or "").strip()
        for cpro, cmun, _control, nombre in sheet.iter_rows(
            min_row=4, values_only=True,
        ):
            province_code = str(cpro or "").zfill(2)
            municipality_code = str(cmun or "").zfill(3)
            municipality_name = str(nombre or "").strip()
            if not municipality_name or province_code not in CCAA_POR_PROVINCIA:
                continue
            rows.append((
                f"{province_code}{municipality_code}",
                CCAA_POR_PROVINCIA[province_code],
                provincia,
                municipality_name,
            ))
    if len(rows) < 8000:
        raise ValueError(f"El fichero solo contiene {len(rows)} municipios")
    destino.parent.mkdir(parents=True, exist_ok=True)
    with destino.open("w", encoding="utf-8", newline="") as output:
        writer = csv.writer(output, lineterminator="\n")
        writer.writerow(("codigo_ine", "ccaa", "provincia", "nombre"))
        writer.writerows(rows)
    return len(rows)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("origen", type=Path)
    parser.add_argument("destino", type=Path)
    args = parser.parse_args()
    print(f"Municipios generados: {generar(args.origen, args.destino)}")


if __name__ == "__main__":
    main()
